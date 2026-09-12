"""Helper xuất file Excel (.xlsx) dùng chung cho các route (openpyxl).

Tách riêng để exam.py / ngan_hang.py / phong_thi.py cùng dùng, tránh lặp code.
"""
import io
from flask import Response


def xuat_xlsx(sheet_title, tieu_de, info_lines, heads, rows, widths, fname):
    """Dựng 1 file .xlsx trong bộ nhớ và trả về Response để tải xuống.

    - sheet_title : tên sheet (cắt 31 ký tự theo giới hạn Excel)
    - tieu_de     : dòng tiêu đề lớn ở đầu
    - info_lines  : danh sách dòng thông tin phụ (vd 'Tổng số câu: 10')
    - heads       : danh sách tên cột (hàng tiêu đề tô nền xanh)
    - rows        : danh sách các hàng dữ liệu (mỗi hàng là list)
    - widths      : danh sách bề rộng từng cột
    - fname       : tên file tải về
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = sheet_title[:31]
    ws.append([tieu_de])
    for line in info_lines:
        ws.append([line])
    ws.append([])
    ws.append(heads)
    hrow = ws.max_row
    fill = PatternFill('solid', fgColor='2563EB')
    font = Font(bold=True, color='FFFFFF')
    for col in range(1, len(heads) + 1):
        cell = ws.cell(row=hrow, column=col)
        cell.fill = fill; cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'})
