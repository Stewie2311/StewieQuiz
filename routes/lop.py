"""
Quản lý LỚP và DANH SÁCH SINH VIÊN.

Giáo viên tạo lớp, thêm sinh viên từng người hoặc nạp cả danh sách từ Excel.
Đây chỉ là sổ danh sách; thí sinh vào phòng thi vẫn qua luồng riêng ở
phong_thi.py, hai bên khớp nhau bằng MSSV.

MÔ HÌNH QUYỀN (dễ nhầm nên ghi rõ):
  - XEM / TẢI EXCEL : mọi giáo viên, kể cả lớp của đồng nghiệp. Cố ý mở như vậy
                      để giáo viên coi thi hộ vẫn tra được danh sách.
  - SỬA / XÓA / THÊM: chỉ người tạo lớp, hoặc admin.
Hai hàm `_lay_lop_de_xem` và `_lay_lop_co_quyen` bên dưới thể hiện đúng hai
mức này — dùng nhầm hàm là mở toang quyền sửa cho người ngoài.

Import Excel đi theo 2 bước: xem trước (`import_preview`, chỉ đọc và soi lỗi)
rồi mới lưu (`import_save`). Nhờ vậy giáo viên thấy dòng nào trùng/thiếu trước
khi ghi vào cơ sở dữ liệu.
"""
import io
from flask import (Blueprint, render_template, request, redirect,
                   session, flash, url_for, jsonify, Response)
from database import get_db_connection
from decorators import giao_vien_required

lop_bp = Blueprint('lop', __name__)


# ============== HÀM DÙNG CHUNG (quyền + làm sạch dữ liệu) ==============

def _lay_lop_co_quyen(cursor, lop_id):
    """Lấy lớp cho thao tác GHI (sửa/xóa/thêm SV): chỉ người tạo hoặc admin.

    Trả về (lop, None) nếu được phép, ngược lại (None, 'notfound'|'forbidden').
    """
    cursor.execute("SELECT * FROM lop WHERE id=%s", (lop_id,))
    lop = cursor.fetchone()
    if not lop:
        return None, 'notfound'
    if session.get('role') != 'admin' and lop['user_id'] != session['user_id']:
        return None, 'forbidden'
    return lop, None


def _lay_lop_de_xem(cursor, lop_id):
    """Lấy lớp cho thao tác ĐỌC (xem danh sách/tải Excel): giáo viên nào cũng được.

    Cố tình KHÔNG kiểm tra chủ sở hữu — xem mô hình quyền ở đầu file.
    """
    cursor.execute("SELECT * FROM lop WHERE id=%s", (lop_id,))
    lop = cursor.fetchone()
    return (lop, None) if lop else (None, 'notfound')


def _la_chu_lop(lop):
    """Người dùng hiện tại có được sửa lớp này không? Template dùng để ẩn/hiện nút."""
    return session.get('role') == 'admin' or lop.get('user_id') == session.get('user_id')


def _mssv_str(v):
    """Ép MSSV về chuỗi sạch.

    Excel lưu ô toàn chữ số thành kiểu số thực, nên MSSV 2001210001 đọc ra là
    2001210001.0. Cứ str() thẳng thì cái đuôi ".0" dính vào mã số, và về sau
    khớp MSSV với thí sinh vào phòng thi sẽ trượt hết.
    """
    if v is None:
        return ''
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return str(v).strip()


# ============== TRANG DANH SÁCH LỚP + TẠO/SỬA/XÓA LỚP ==============

@lop_bp.route('/lop')
@giao_vien_required
def danh_sach_lop():
    """Trang /lop: liệt kê mọi lớp trong trường, kèm số liệu cho panel thống kê.

    Truy vấn lấy TẤT CẢ lớp, sau đó tách ra hai nhóm ("lớp của tôi" và "tất cả")
    cùng vài con số sĩ số cho giao diện. Tính sẵn ở server để template khỏi
    phải lặp và đếm.
    """
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    cur.execute("""
        SELECT l.*, u.ho_ten AS nguoi_tao,
               (SELECT COUNT(*) FROM lop_sinh_vien sv WHERE sv.lop_id=l.id) AS si_so
        FROM lop l
        LEFT JOIN nguoi_dung u ON u.id = l.user_id
        ORDER BY l.id DESC
    """)
    lops = cur.fetchall()
    cur.close(); conn.close()

    # Phân loại lớp theo sĩ số: nhỏ (<20), vừa (20–50), lớn (>50).
    all_lops = lops
    my_lops  = [l for l in lops if l['user_id'] == session['user_id']]
    my_lop_nho  = [l for l in my_lops if (l['si_so'] or 0) < 20]
    my_lop_vua  = [l for l in my_lops if 20 <= (l['si_so'] or 0) <= 50]
    my_lop_lon  = [l for l in my_lops if (l['si_so'] or 0) > 50]
    my_tong_si_so = sum(l['si_so'] or 0 for l in my_lops)
    my_so_lop = len(my_lops)

    tong_si_so = sum(l['si_so'] or 0 for l in all_lops)
    lop_nho  = [l for l in all_lops if (l['si_so'] or 0) < 20]
    lop_vua = [l for l in all_lops if 20 <= (l['si_so'] or 0) <= 50]
    lop_lon  = [l for l in all_lops if (l['si_so'] or 0) > 50]

    return render_template('lop_hoc.html', lops=all_lops,
                           my_lops=my_lops,
                           all_tong_si_so=tong_si_so,
                           all_lop_nho=lop_nho, all_lop_vua=lop_vua, all_lop_lon=lop_lon,
                           my_tong_si_so=my_tong_si_so, my_so_lop=my_so_lop,
                           my_lop_nho=my_lop_nho, my_lop_vua=my_lop_vua, my_lop_lon=my_lop_lon)


@lop_bp.route('/lop/tao', methods=['POST'])
@giao_vien_required
def tao_lop():
    """Tạo lớp rỗng; người bấm nút trở thành chủ lớp."""
    ten = (request.form.get('ten_lop') or '').strip()
    if not ten:
        flash('Vui lòng nhập tên lớp!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    conn = get_db_connection(); cur = conn.cursor()
    cur.execute("INSERT INTO lop (ten_lop, user_id) VALUES (%s, %s)",
                (ten, session['user_id']))
    conn.commit(); cur.close(); conn.close()
    flash('Đã tạo lớp mới!', 'success')
    return redirect(url_for('lop.danh_sach_lop'))


@lop_bp.route('/lop/<int:lop_id>/sua', methods=['POST'])
@giao_vien_required
def sua_lop(lop_id):
    """Đổi tên lớp — chỉ người tạo (hoặc admin)."""
    ten = (request.form.get('ten_lop') or '').strip()
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_co_quyen(cur, lop_id)
    if err:
        cur.close(); conn.close()
        flash('Không tìm thấy lớp!' if err == 'notfound' else 'Bạn không phải người tạo lớp này!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    if not ten:
        cur.close(); conn.close()
        flash('Vui lòng nhập tên lớp!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    cur.execute("UPDATE lop SET ten_lop=%s WHERE id=%s", (ten, lop_id))
    conn.commit(); cur.close(); conn.close()
    flash('Đã cập nhật tên lớp.', 'success')
    return redirect(url_for('lop.danh_sach_lop'))


@lop_bp.route('/lop/<int:lop_id>/xoa', methods=['POST'])
@giao_vien_required
def xoa_lop(lop_id):
    """Xóa lớp; sinh viên trong lớp bị xóa theo (ràng buộc khóa ngoại)."""
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_co_quyen(cur, lop_id)
    if err:
        cur.close(); conn.close()
        flash('Không tìm thấy lớp!' if err == 'notfound' else 'Không có quyền!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    cur.execute("DELETE FROM lop WHERE id=%s", (lop_id,))
    conn.commit(); cur.close(); conn.close()
    flash('Đã xóa lớp.', 'success')
    return redirect(url_for('lop.danh_sach_lop'))


# ============== CHI TIẾT LỚP: DANH SÁCH SINH VIÊN, XUẤT EXCEL, THÊM/XÓA SV ==============

@lop_bp.route('/lop/<int:lop_id>')
@giao_vien_required
def chi_tiet_lop(lop_id):
    """Trang danh sách sinh viên của một lớp.

    Cờ `is_owner` gửi xuống template để ẩn các nút Thêm/Xóa/Sửa với giáo viên
    chỉ ghé xem lớp của người khác.
    """
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_de_xem(cur, lop_id)
    if err:
        cur.close(); conn.close()
        flash('Không tìm thấy lớp!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    cur.execute("SELECT * FROM lop_sinh_vien WHERE lop_id=%s ORDER BY ma_so_sv", (lop_id,))
    svs = cur.fetchall()
    cur.close(); conn.close()
    return render_template('lop_chi_tiet.html', lop=lop, svs=svs, is_owner=_la_chu_lop(lop))


@lop_bp.route('/lop/<int:lop_id>/excel')
@giao_vien_required
def xuat_excel(lop_id):
    """Tải danh sách lớp về dưới dạng .xlsx (mọi giáo viên đều tải được)."""
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_de_xem(cur, lop_id)
    if err:
        cur.close(); conn.close()
        flash('Không tìm thấy lớp!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    cur.execute("SELECT * FROM lop_sinh_vien WHERE lop_id=%s ORDER BY ma_so_sv", (lop_id,))
    svs = cur.fetchall()
    cur.close(); conn.close()

    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'DanhSach'
    ws.append(['STT', 'MSSV', 'Họ và tên', 'Email'])
    for i, s in enumerate(svs, start=1):
        ws.append([i, s['ma_so_sv'], s['ho_ten'], s.get('email') or ''])
    for idx, w in enumerate([6, 20, 30, 30], start=1):
        ws.column_dimensions[chr(64 + idx)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)

    # Tên lớp có thể chứa dấu gạch chéo, dấu hai chấm... — nhét thẳng vào tên
    # file thì trình duyệt hiểu sai hoặc hỏng đường dẫn, nên thay hết bằng '_'.
    # Phần dấu tiếng Việt được giữ nhờ filename*=UTF-8'' ở header bên dưới.
    safe = ''.join(c if c.isalnum() or c in ' -_' else '_' for c in (lop['ten_lop'] or 'lop')).strip() or 'lop'
    from urllib.parse import quote
    fname = f'DanhSach_{safe}.xlsx'
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(fname)}"})


@lop_bp.route('/lop/<int:lop_id>/them_sv', methods=['POST'])
@giao_vien_required
def them_sv(lop_id):
    """Thêm một sinh viên vào lớp.

    MSSV có ràng buộc UNIQUE theo lớp, nên thêm trùng sẽ ném lỗi ở INSERT —
    bắt lại và báo "đã tồn tại" thay vì bung lỗi 500 ra người dùng.
    """
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_co_quyen(cur, lop_id)
    if err:
        cur.close(); conn.close()
        flash('Không có quyền!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    mssv  = (request.form.get('ma_so_sv') or '').strip()
    hoten = (request.form.get('ho_ten') or '').strip()
    email = (request.form.get('email') or '').strip()
    if not mssv or not hoten:
        flash('Cần nhập MSSV và Họ tên!', 'danger')
        cur.close(); conn.close()
        return redirect(url_for('lop.chi_tiet_lop', lop_id=lop_id))
    try:
        cur.execute("""INSERT INTO lop_sinh_vien (lop_id, ma_so_sv, ho_ten, email)
                       VALUES (%s,%s,%s,%s)""", (lop_id, mssv, hoten, email or None))
        conn.commit()
        flash('Đã thêm sinh viên.', 'success')
    except Exception:
        conn.rollback()
        flash(f'MSSV "{mssv}" đã tồn tại trong lớp này.', 'danger')
    cur.close(); conn.close()
    return redirect(url_for('lop.chi_tiet_lop', lop_id=lop_id))


@lop_bp.route('/lop/<int:lop_id>/xoa_sv/<int:sv_id>', methods=['POST'])
@giao_vien_required
def xoa_sv(lop_id, sv_id):
    """Xóa một sinh viên khỏi lớp.

    Câu DELETE có kèm `AND lop_id=%s`: không có nó, người ta đoán được id sinh
    viên bất kỳ và xóa người của lớp khác qua đường lớp mình sở hữu.
    """
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_co_quyen(cur, lop_id)
    if err:
        cur.close(); conn.close()
        flash('Không có quyền!', 'danger')
        return redirect(url_for('lop.danh_sach_lop'))
    cur.execute("DELETE FROM lop_sinh_vien WHERE id=%s AND lop_id=%s", (sv_id, lop_id))
    conn.commit(); cur.close(); conn.close()
    flash('Đã xóa sinh viên khỏi lớp.', 'success')
    return redirect(url_for('lop.chi_tiet_lop', lop_id=lop_id))


# ============== NHẬP DANH SÁCH TỪ EXCEL (xem trước -> lưu) ==============

@lop_bp.route('/lop/mau_excel')
@giao_vien_required
def mau_excel():
    """Tải file Excel mẫu, có sẵn hàng tiêu đề và 2 dòng ví dụ để làm theo."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'DanhSach'
    ws.append(['MSSV', 'Họ và tên', 'Email'])
    ws.append(['2001210001', 'Nguyễn Văn A', 'a@example.com'])
    ws.append(['2001210002', 'Trần Thị B', ''])
    for i, w in enumerate([20, 30, 30], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=mau_danh_sach_lop.xlsx'})


@lop_bp.route('/lop/<int:lop_id>/import_preview', methods=['POST'])
@giao_vien_required
def import_preview(lop_id):
    """BƯỚC 1 của import: đọc file Excel, soi lỗi từng dòng, KHÔNG ghi gì vào DB.

    Trả về danh sách dòng kèm nhãn lỗi (thiếu MSSV / thiếu họ tên / đã có trong
    lớp / trùng ngay trong file) để giáo viên nhìn thấy trước khi quyết định lưu.

    Cột trong file không cần đúng thứ tự: hàm `tim` bên dưới dò tiêu đề theo từ
    khóa, nên "MSSV", "Mã số SV", "Student ID" đều nhận ra được.
    """
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_co_quyen(cur, lop_id)
    if err:
        cur.close(); conn.close()
        return jsonify({'ok': False, 'error': 'Không có quyền'}), 403
    cur.execute("SELECT ma_so_sv FROM lop_sinh_vien WHERE lop_id=%s", (lop_id,))
    da_co = {r['ma_so_sv'] for r in cur.fetchall()}
    cur.close(); conn.close()

    f = request.files.get('file')
    if not f or not f.filename.lower().endswith('.xlsx'):
        return jsonify({'ok': False, 'error': 'Vui lòng chọn file Excel (.xlsx)'}), 400

    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return jsonify({'ok': False, 'error': 'Không đọc được file. Hãy dùng đúng file mẫu .xlsx'}), 400

    if not rows:
        return jsonify({'ok': False, 'error': 'File trống'}), 400

    header = [str(c).strip().lower() if c is not None else '' for c in rows[0]]

    def tim(*names):
        """Tìm chỉ số cột có tiêu đề chứa một trong các từ khóa; -1 nếu không có."""
        for i, h in enumerate(header):
            if any(n in h for n in names):
                return i
        return -1

    i_mssv  = tim('mssv', 'mã số', 'ma so', 'student')
    i_ten   = tim('họ', 'ho ten', 'tên', 'ten', 'name')
    i_email = tim('email', 'mail')
    if i_mssv < 0 or i_ten < 0:
        return jsonify({'ok': False,
            'error': 'Không tìm thấy cột MSSV hoặc Họ tên. Hãy dùng file mẫu (hàng đầu là tiêu đề).'}), 400

    seen = set(); preview = []
    for r in rows[1:]:
        mssv  = _mssv_str(r[i_mssv] if i_mssv < len(r) else None)
        ten   = (str(r[i_ten]).strip() if i_ten < len(r) and r[i_ten] is not None else '')
        email = (str(r[i_email]).strip() if 0 <= i_email < len(r) and r[i_email] is not None else '')
        if not mssv and not ten:
            continue
        if not mssv:
            loi = 'Thiếu MSSV'
        elif not ten:
            loi = 'Thiếu họ tên'
        elif mssv in da_co:
            loi = 'Đã có trong lớp'
        elif mssv in seen:
            loi = 'Trùng trong file'
        else:
            loi = ''
            seen.add(mssv)
        preview.append({'mssv': mssv, 'ho_ten': ten, 'email': email,
                        'loi': loi, 'ok': loi == ''})

    return jsonify({'ok': True, 'rows': preview,
                    'hop_le': sum(1 for x in preview if x['ok'])})


@lop_bp.route('/lop/<int:lop_id>/import_save', methods=['POST'])
@giao_vien_required
def import_save(lop_id):
    """BƯỚC 2 của import: lưu các dòng giáo viên đã duyệt ở màn hình xem trước.

    Dùng INSERT IGNORE nên MSSV trùng bị bỏ qua lặng lẽ thay vì làm hỏng cả mẻ.
    `added` cộng dồn từ rowcount nên đếm đúng số dòng THỰC SỰ được thêm.
    """
    conn = get_db_connection(); cur = conn.cursor(dictionary=True)
    lop, err = _lay_lop_co_quyen(cur, lop_id)
    if err:
        cur.close(); conn.close()
        return jsonify({'ok': False}), 403
    data = request.get_json(silent=True) or {}
    rows = data.get('rows') or []
    cur2 = conn.cursor()
    added = 0
    for r in rows:
        mssv  = (r.get('mssv') or '').strip()
        ten   = (r.get('ho_ten') or '').strip()
        email = (r.get('email') or '').strip()
        if not mssv or not ten:
            continue
        try:
            cur2.execute("""INSERT IGNORE INTO lop_sinh_vien (lop_id, ma_so_sv, ho_ten, email)
                            VALUES (%s,%s,%s,%s)""", (lop_id, mssv, ten, email or None))
            added += cur2.rowcount
        except Exception:
            pass
    conn.commit()
    cur2.close(); cur.close(); conn.close()
    return jsonify({'ok': True, 'added': added})
